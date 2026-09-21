"""Turn raw credential/job files into text, tagged with where each piece came from.

Provenance (the source filename) travels with every piece of text so the
credential profile step (and the fidelity check later) can point back to the
original document. Every failure is per-file/per-link: one bad file never
stops the rest of the batch.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from resume_match.llm import LLMConfig, LLMError, complete

TEXT_EXTENSIONS = {".txt", ".md"}
DOCX_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}

MAX_IMAGE_BYTES = 8 * 1024 * 1024  # per-file size cap; larger scans are skipped with a notice


@dataclass
class Source:
    """One ingested document or link, ready to feed into a profile/job-analysis prompt."""

    name: str  # filename or URL, used as the provenance tag
    text: str
    kind: str  # "text" | "image" | "link"
    error: str | None = None


@dataclass
class IngestResult:
    sources: list[Source] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def combined_text(self) -> str:
        """All source text concatenated with provenance headers, for prompting."""
        parts = []
        for source in self.sources:
            if source.text:
                parts.append(f"--- Source: {source.name} ---\n{source.text}")
        return "\n\n".join(parts)


def ingest_files(files: list[tuple[str, bytes]], llm_config: LLMConfig | None = None) -> IngestResult:
    """Ingest a list of (filename, raw_bytes) pairs into an IngestResult.

    `llm_config` is required only if the batch includes images (vision
    transcription); text-only batches work without it.
    """
    result = IngestResult()
    for filename, data in files:
        try:
            source = _ingest_one(filename, data, llm_config)
            if source is not None:
                result.sources.append(source)
        except Exception as exc:  # noqa: BLE001 - deliberately broad, one bad file must not sink the batch
            result.warnings.append(f"Could not read {filename}: {exc}")
    return result


def ingest_links(urls: list[str]) -> IngestResult:
    """Fetch each URL and strip it to text. Each link fails independently."""
    result = IngestResult()
    for url in urls:
        url = url.strip()
        if not url:
            continue
        try:
            text = _fetch_link_text(url)
            result.sources.append(Source(name=url, text=text, kind="link"))
        except Exception as exc:  # noqa: BLE001
            result.warnings.append(
                f"Could not fetch {url}: {exc}. If this is LinkedIn or a dynamic page, "
                "save it as a PDF instead (Print -> Save as PDF) and add the file to the folder."
            )
    return result


def gather_credentials(
    files: list[tuple[str, bytes]], llm_config: LLMConfig | None = None
) -> tuple[IngestResult, str]:
    """Ingest a credentials folder/upload set, handling `links.txt` and
    `preferences.txt`/`.md` specially per PRD §5.1.

    Returns (result, preferences_text). `links.txt` is read as one URL per
    line and fetched rather than treated as document text; a preferences
    file is pulled out and returned separately rather than folded into the
    profile source text.
    """
    ordinary_files = []
    link_urls: list[str] = []
    preferences_text = ""

    for filename, data in files:
        base = filename.rsplit("/", 1)[-1].lower()
        if base == "links.txt":
            link_urls.extend(_decode_text(data).splitlines())
        elif base in ("preferences.txt", "preferences.md"):
            preferences_text = _decode_text(data)
        else:
            ordinary_files.append((filename, data))

    result = ingest_files(ordinary_files, llm_config)
    if link_urls:
        link_result = ingest_links(link_urls)
        result.sources.extend(link_result.sources)
        result.warnings.extend(link_result.warnings)

    return result, preferences_text


def identify_resume_candidates(sources: list[Source]) -> list[str]:
    """Filenames that look like a resume/CV, for the base-resume selector."""
    return [s.name for s in sources if "resume" in s.name.lower() or "cv" in s.name.lower()]


def _ingest_one(filename: str, data: bytes, llm_config: LLMConfig | None) -> Source | None:
    ext = _extension(filename)

    if ext in TEXT_EXTENSIONS:
        return Source(name=filename, text=_decode_text(data), kind="text")

    if ext in DOCX_EXTENSIONS:
        return Source(name=filename, text=_extract_docx_text(data), kind="text")

    if ext in PDF_EXTENSIONS:
        return Source(name=filename, text=_extract_pdf_text(data), kind="text")

    if ext in IMAGE_EXTENSIONS:
        if len(data) > MAX_IMAGE_BYTES:
            raise ValueError(f"image is larger than {MAX_IMAGE_BYTES // (1024 * 1024)} MB; skipped")
        if llm_config is None:
            raise ValueError("no model configured to transcribe images")
        return Source(name=filename, text=_transcribe_image(data, llm_config), kind="image")

    raise ValueError(f"unsupported file type '{ext}'")


def _extension(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def _decode_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _extract_docx_text(data: bytes) -> str:
    import io

    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _extract_pdf_text(data: bytes) -> str:
    import io

    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _transcribe_image(data: bytes, llm_config: LLMConfig) -> str:
    prompt = (
        "Transcribe every piece of text visible in this image exactly as written. "
        "This is a scanned academic or professional document (diploma, transcript, "
        "certificate). Preserve names, dates, numbers, and titles precisely. "
        "Do not summarize, interpret, or add anything not visible in the image."
    )
    try:
        return complete(llm_config, prompt, images=[data])
    except LLMError as exc:
        raise ValueError(str(exc)) from exc


def _fetch_link_text(url: str) -> str:
    import requests
    from bs4 import BeautifulSoup

    response = requests.get(
        url,
        timeout=15,
        headers={"User-Agent": "Mozilla/5.0 (compatible; SenSymResumeMatch/1.0)"},
    )
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)
